from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook


YEARS = [2010, 2015, 2020]
PERIODS = [(2010, 2015), (2015, 2020), (2010, 2020)]
SCHEMA_VERSION = "1.1.0"
METHOD_VERSION = "2026-09-03"
TOLERANCE = 1e-8
APP_ROOT = Path(__file__).resolve().parents[1]
RESULTS_RELATIVE_PATH = Path("projects") / "interim-presentation" / "코드" / "결과" / "주요결과"
RESULTS_SOURCE_PATH = RESULTS_RELATIVE_PATH.as_posix()


class ExportError(RuntimeError):
    pass


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, column).value).strip(): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(1, column).value is not None
    }


def _require_headers(ws, expected: list[str]) -> dict[str, int]:
    headers = _headers(ws)
    missing = [name for name in expected if name not in headers]
    if missing:
        raise ExportError(f"{ws.title}: missing columns: {', '.join(missing)}")
    return headers


def _number(value: Any, location: str) -> float:
    if value is None or isinstance(value, bool):
        raise ExportError(f"missing numeric value at {location}")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ExportError(f"non-numeric value at {location}: {value!r}") from exc
    if not math.isfinite(result):
        raise ExportError(f"non-finite value at {location}")
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_necessary_labour(path: Path, generated_at: str) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExportError(f"cannot open workbook: {path}") from exc
    try:
        if "summary" not in workbook.sheetnames:
            raise ExportError("05_constant_value.xlsx: missing sheet: summary")
        ws = workbook["summary"]
        headers = _require_headers(ws, ["year", "necessary_labour_constant"])
        rows: list[dict[str, Any]] = []
        seen: set[int] = set()
        for row in range(2, ws.max_row + 1):
            raw_year = ws.cell(row, headers["year"]).value
            if raw_year is None:
                continue
            year = int(raw_year)
            if year in seen:
                raise ExportError(f"summary: duplicate year {year}")
            seen.add(year)
            if year in YEARS:
                value_cell = ws.cell(row, headers["necessary_labour_constant"])
                rows.append({"year": year, "value": _number(value_cell.value, f"summary!{value_cell.coordinate}"), "sourceCell": value_cell.coordinate})
        if [item["year"] for item in rows] != YEARS:
            raise ExportError(f"summary years must be exactly {YEARS} in order")
        if headers["year"] != 1 or headers["necessary_labour_constant"] != 11:
            raise ExportError("05 summary layout changed; expected year=A and necessary_labour_constant=K")
        return {
            "schemaVersion": SCHEMA_VERSION,
            "generatedAt": generated_at,
            "unit": "hours",
            "source": {"workbook": f"{RESULTS_SOURCE_PATH}/05_constant_value.xlsx", "sheet": "summary", "yearRange": "A2:A4", "valueRange": "K2:K4"},
            "methodology": {"id": "necessary-labour-constant-price", "version": METHOD_VERSION, "priceBasis": "2020 constant prices", "sectorClassification": "K77"},
            "series": rows,
        }
    finally:
        workbook.close()


def read_decomposition(path: Path, generated_at: str) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExportError(f"cannot open workbook: {path}") from exc
    try:
        required_sheets = ["summary", *(f"{start}-{end}" for start, end in PERIODS)]
        missing_sheets = [sheet for sheet in required_sheets if sheet not in workbook.sheetnames]
        if missing_sheets:
            raise ExportError(f"06_decomposition.xlsx: missing sheets: {', '.join(missing_sheets)}")
        summary = workbook["summary"]
        headers = _require_headers(summary, ["period", "start_year", "end_year", "V_start", "V_end", "delta_V", "b_effect", "t_effect", "check"])
        expected_columns = {"period": 1, "start_year": 2, "end_year": 3, "V_start": 4, "V_end": 5, "delta_V": 6, "b_effect": 7, "t_effect": 8, "check": 9}
        if any(headers[name] != column for name, column in expected_columns.items()):
            raise ExportError("06 summary layout changed; expected columns A:I")
        periods: list[dict[str, Any]] = []
        for row, (expected_start, expected_end) in enumerate(PERIODS, start=2):
            start = int(summary.cell(row, headers["start_year"]).value)
            end = int(summary.cell(row, headers["end_year"]).value)
            period = str(summary.cell(row, headers["period"]).value)
            if (start, end) != (expected_start, expected_end) or period != f"{start}-{end}":
                raise ExportError(f"summary!A{row}: unexpected period ordering")
            contribution_ws = workbook[period]
            contribution_headers = _require_headers(contribution_ws, ["code", "name", "delta_V_sector", "b_effect_sector", "t_effect_sector"])
            if [contribution_headers[key] for key in ["delta_V_sector", "b_effect_sector", "t_effect_sector"]] != [13, 14, 15]:
                raise ExportError(f"{period}: contribution columns must be M:O")
            contributions: list[dict[str, Any]] = []
            for sector_row in range(2, contribution_ws.max_row + 1):
                raw_code = contribution_ws.cell(sector_row, contribution_headers["code"]).value
                if raw_code is None:
                    continue
                raw_name = contribution_ws.cell(sector_row, contribution_headers["name"]).value
                if raw_name is None or not str(raw_name).strip():
                    raise ExportError(f"{period}!B{sector_row}: missing sector name")
                contributions.append({
                    "code": str(raw_code),
                    "name": str(raw_name),
                    "totalChange": _number(contribution_ws.cell(sector_row, contribution_headers["delta_V_sector"]).value, f"{period}!M{sector_row}"),
                    "basketEffect": _number(contribution_ws.cell(sector_row, contribution_headers["b_effect_sector"]).value, f"{period}!N{sector_row}"),
                    "embodiedLabourEffect": _number(contribution_ws.cell(sector_row, contribution_headers["t_effect_sector"]).value, f"{period}!O{sector_row}"),
                    "sourceRow": sector_row,
                })
            if len(contributions) != 77:
                raise ExportError(f"{period}: expected 77 sector rows, found {len(contributions)}")
            periods.append({
                "period": period,
                "fromYear": start,
                "toYear": end,
                "startValue": _number(summary.cell(row, headers["V_start"]).value, f"summary!D{row}"),
                "endValue": _number(summary.cell(row, headers["V_end"]).value, f"summary!E{row}"),
                "totalChange": _number(summary.cell(row, headers["delta_V"]).value, f"summary!F{row}"),
                "basketEffect": _number(summary.cell(row, headers["b_effect"]).value, f"summary!G{row}"),
                "embodiedLabourEffect": _number(summary.cell(row, headers["t_effect"]).value, f"summary!H{row}"),
                "residual": _number(summary.cell(row, headers["check"]).value, f"summary!I{row}"),
                "sourceRow": row,
                "contributions": contributions,
            })
        return {
            "schemaVersion": SCHEMA_VERSION,
            "generatedAt": generated_at,
            "unit": "hours",
            "source": {"workbook": f"{RESULTS_SOURCE_PATH}/06_decomposition.xlsx", "summarySheet": "summary", "summaryRange": "A1:I4", "contributionColumns": "M:O"},
            "methodology": {"id": "symmetric-two-factor-decomposition", "version": METHOD_VERSION, "priceBasis": "2020 constant prices", "sectorClassification": "K77"},
            "periods": periods,
        }
    finally:
        workbook.close()


def build_validation(levels: dict[str, Any], decomposition: dict[str, Any], level_path: Path, decomposition_path: Path, generated_at: str) -> dict[str, Any]:
    level_by_year = {item["year"]: item["value"] for item in levels["series"]}
    checks = [
        {"id": "required-cells", "status": "pass", "message": "Required headers and numeric cells are present."},
        {"id": "year-consistency", "status": "pass", "message": "Years are exactly 2010, 2015, and 2020 in canonical order."},
        {"id": "sector-coverage", "status": "pass", "message": "Each decomposition sheet contains 77 identified sector rows."},
    ]
    max_identity_residual = 0.0
    max_sector_residual = 0.0
    for period in decomposition["periods"]:
        start, end = period["fromYear"], period["toYear"]
        endpoint_error = max(abs(period["startValue"] - level_by_year[start]), abs(period["endValue"] - level_by_year[end]))
        if endpoint_error > TOLERANCE:
            raise ExportError(f"{period['period']}: decomposition endpoints differ from 05 summary by {endpoint_error}")
        identity = period["basketEffect"] + period["embodiedLabourEffect"] - period["totalChange"]
        max_identity_residual = max(max_identity_residual, abs(identity), abs(period["residual"]))
        for key in ["totalChange", "basketEffect", "embodiedLabourEffect"]:
            sector_sum = sum(item[key] for item in period["contributions"])
            max_sector_residual = max(max_sector_residual, abs(sector_sum - period[key]))
    if max_identity_residual > TOLERANCE:
        raise ExportError(f"decomposition identity residual exceeds tolerance: {max_identity_residual}")
    if max_sector_residual > TOLERANCE:
        raise ExportError(f"sector contributions do not reconcile: {max_sector_residual}")
    checks.extend([
        {"id": "workbook-endpoints", "status": "pass", "message": "Decomposition endpoints match necessary-labour levels within 1e-8 hours."},
        {"id": "decomposition-identity", "status": "pass", "message": f"Maximum total identity residual is {max_identity_residual:.3e} hours."},
        {"id": "sector-reconciliation", "status": "pass", "message": f"Maximum sector-to-summary residual is {max_sector_residual:.3e} hours."},
    ])
    return {
        "schemaVersion": SCHEMA_VERSION,
        "generatedAt": generated_at,
        "scope": "exporter-only",
        "status": "pass",
        "sourceWorkbooks": [
            {"path": f"{RESULTS_SOURCE_PATH}/05_constant_value.xlsx", "sha256": _sha256(level_path)},
            {"path": f"{RESULTS_SOURCE_PATH}/06_decomposition.xlsx", "sha256": _sha256(decomposition_path)},
        ],
        "checks": checks,
        "limitation": "These checks validate extraction, schema, cross-workbook endpoints, and saved decomposition identities. They do not rerun or independently validate the original calculation pipeline.",
    }


def validate_documents(documents: dict[str, dict[str, Any]], schema_dir: Path) -> None:
    format_checker = FormatChecker()
    for name, document in documents.items():
        schema_path = schema_dir / f"{name}.schema.json"
        try:
            schema = json.loads(schema_path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise ExportError(f"cannot read schema: {schema_path}") from exc
        validator = Draft202012Validator(schema, format_checker=format_checker)
        errors = sorted(validator.iter_errors(document), key=lambda item: list(item.path))
        if errors:
            detail = "; ".join(f"{'.'.join(map(str, error.path)) or '<root>'}: {error.message}" for error in errors[:5])
            raise ExportError(f"schema validation failed for {name}: {detail}")


def write_documents(documents: dict[str, dict[str, Any]], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    temporary_paths: list[tuple[Path, Path]] = []
    try:
        for name, document in documents.items():
            handle, temporary_name = tempfile.mkstemp(prefix=f".{name}.", suffix=".tmp", dir=output_dir)
            temporary_path = Path(temporary_name)
            with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
                json.dump(document, stream, ensure_ascii=False, indent=2)
                stream.write("\n")
            temporary_paths.append((temporary_path, output_dir / f"{name}.json"))
        for temporary_path, destination in temporary_paths:
            os.replace(temporary_path, destination)
    finally:
        for temporary_path, _ in temporary_paths:
            if temporary_path.exists():
                temporary_path.unlink()


def export_results(level_path: Path, decomposition_path: Path, schema_dir: Path, output_dir: Path) -> dict[str, dict[str, Any]]:
    if not level_path.is_file():
        raise ExportError(f"missing workbook: {level_path}")
    if not decomposition_path.is_file():
        raise ExportError(f"missing workbook: {decomposition_path}")
    generated_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    levels = read_necessary_labour(level_path, generated_at)
    decomposition = read_decomposition(decomposition_path, generated_at)
    validation = build_validation(levels, decomposition, level_path, decomposition_path, generated_at)
    documents = {"necessary_labour": levels, "decomposition": decomposition, "validation": validation}
    validate_documents(documents, schema_dir)
    write_documents(documents, output_dir)
    return documents


def resolve_repository_root(configured_root: Path | None) -> Path:
    """Return the research-data checkout used for workbook input and JSON output."""
    value = configured_root or (Path(os.environ["LOCAL_REPOSITORY_ROOT"]) if os.environ.get("LOCAL_REPOSITORY_ROOT") else None)
    if value is None:
        raise ExportError("set LOCAL_REPOSITORY_ROOT or pass --repository-root to the Obsidian-Vault checkout")
    root = value.expanduser().resolve()
    if not root.is_dir():
        raise ExportError(f"repository root is not a directory: {root}")
    return root


def main() -> int:
    parser = argparse.ArgumentParser(description="Export canonical Calc workbooks to validated dashboard JSON.")
    parser.add_argument("--repository-root", type=Path, help="Absolute path to the Obsidian-Vault research-data checkout.")
    parser.add_argument("--results-dir", type=Path, help="Override the canonical workbook directory under the research-data checkout.")
    parser.add_argument("--schema-dir", type=Path, default=APP_ROOT / "schemas")
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    try:
        repository_root = resolve_repository_root(args.repository_root)
        results_dir = args.results_dir or repository_root / RESULTS_RELATIVE_PATH
        output_dir = args.output_dir or results_dir / "dashboard"
        documents = export_results(results_dir / "05_constant_value.xlsx", results_dir / "06_decomposition.xlsx", args.schema_dir, output_dir)
    except ExportError as exc:
        print(f"Export failed: {exc}")
        return 1
    print(f"Exported {len(documents)} validated documents to {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
