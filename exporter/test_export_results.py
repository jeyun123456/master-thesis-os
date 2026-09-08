import json
import os
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from jsonschema import ValidationError, validate
from openpyxl import Workbook

from exporter.export_results import ExportError, export_results, resolve_repository_root


class ExportResultsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.results = self.root / "results"
        self.results.mkdir()
        self.output = self.results / "dashboard"
        self.schemas = Path(__file__).resolve().parents[1] / "schemas"
        self._write_level_fixture()
        self._write_decomposition_fixture()

    def tearDown(self):
        self.temp.cleanup()

    def _write_level_fixture(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "summary"
        ws.append(["year", *[f"unused_{index}" for index in range(2, 11)], "necessary_labour_constant"])
        for year, value in [(2010, 10.0), (2015, 12.0), (2020, 11.0)]:
            ws.append([year, *([0] * 9), value])
        wb.save(self.results / "05_constant_value.xlsx")
        wb.close()

    def _write_decomposition_fixture(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "summary"
        ws.append(["period", "start_year", "end_year", "V_start", "V_end", "delta_V", "b_effect", "t_effect", "check"])
        rows = [
            (2010, 2015, 10.0, 12.0, 2.0, 3.0, -1.0),
            (2015, 2020, 12.0, 11.0, -1.0, 2.0, -3.0),
            (2010, 2020, 10.0, 11.0, 1.0, 5.0, -4.0),
        ]
        for start, end, v0, v1, delta, basket, embodied in rows:
            period = f"{start}-{end}"
            ws.append([period, start, end, v0, v1, delta, basket, embodied, basket + embodied - delta])
            detail = wb.create_sheet(period)
            detail.append(["code", "name", *[f"unused_{index}" for index in range(3, 13)], "delta_V_sector", "b_effect_sector", "t_effect_sector"])
            for index in range(77):
                detail.append([str(index + 1), f"Sector {index + 1}", *([0] * 10), delta / 77, basket / 77, embodied / 77])
        wb.save(self.results / "06_decomposition.xlsx")
        wb.close()

    def test_exports_schema_valid_documents(self):
        documents = export_results(self.results / "05_constant_value.xlsx", self.results / "06_decomposition.xlsx", self.schemas, self.output)
        self.assertEqual([item["value"] for item in documents["necessary_labour"]["series"]], [10.0, 12.0, 11.0])
        self.assertEqual(len(documents["decomposition"]["periods"][0]["contributions"]), 77)
        self.assertEqual(documents["validation"]["scope"], "exporter-only")
        self.assertEqual({path.name for path in self.output.glob("*.json")}, {"necessary_labour.json", "decomposition.json", "validation.json"})

    def test_missing_workbook_does_not_create_output(self):
        with self.assertRaises(ExportError):
            export_results(self.results / "missing.xlsx", self.results / "06_decomposition.xlsx", self.schemas, self.output)
        self.assertFalse(self.output.exists())

    def test_missing_required_header_does_not_replace_existing_output(self):
        self.output.mkdir()
        sentinel = self.output / "necessary_labour.json"
        sentinel.write_text('{"sentinel": true}', encoding="utf-8")
        wb = Workbook()
        wb.active.title = "summary"
        wb.active.append(["year", "wrong_value"])
        wb.save(self.results / "05_constant_value.xlsx")
        wb.close()
        with self.assertRaises(ExportError):
            export_results(self.results / "05_constant_value.xlsx", self.results / "06_decomposition.xlsx", self.schemas, self.output)
        self.assertEqual(json.loads(sentinel.read_text(encoding="utf-8")), {"sentinel": True})

    def test_schema_rejects_missing_provenance(self):
        schema = json.loads((self.schemas / "necessary_labour.schema.json").read_text(encoding="utf-8"))
        with self.assertRaises(ValidationError):
            validate({"schemaVersion": "1.1.0", "unit": "hours", "series": []}, schema)

    def test_repository_root_is_explicit_and_must_contain_calc(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(ExportError):
                resolve_repository_root(None)
        with self.assertRaises(ExportError):
            resolve_repository_root(self.root)
        (self.root / "Calc").mkdir()
        self.assertEqual(resolve_repository_root(self.root), self.root.resolve())


if __name__ == "__main__":
    unittest.main()
